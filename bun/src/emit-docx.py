#!/usr/bin/env python3
"""Pillar 3 flow DOCX — TYPESETTING.md canon precision (Newsreader, margins, rhythm, tables)."""

from __future__ import annotations

import json
import shutil
import sys
import uuid
import zipfile
from pathlib import Path
from lxml import etree as ET

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt

W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
R_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
CT_NS = "http://schemas.openxmlformats.org/package/2006/content-types"
REL_NS = "http://schemas.openxmlformats.org/package/2006/relationships"

NNBSP_PIPE = "\u202f|\u202f"


def load_json(path: Path) -> dict:
    return json.loads(path.read_text(encoding="utf-8"))


def mm_to_pt(mm: float) -> float:
    return mm * 72 / 25.4


def set_run_font(
    run,
    family: str,
    size_pt: float,
    *,
    bold: bool = False,
    italic: bool = False,
    spacing_pt: float = 0,
    lining_nums: bool = False,
    tabular_nums: bool = False,
    oldstyle_nums: bool = False,
    proportional_nums: bool = False,
) -> None:
    run.font.name = family
    run.font.size = Pt(size_pt)
    run.font.bold = bold
    run.font.italic = italic
    rpr = run._element.get_or_add_rPr()
    if spacing_pt:
        spacing = OxmlElement("w:spacing")
        spacing.set(qn("w:val"), str(round(spacing_pt * 20)))
        rpr.append(spacing)
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    rfonts.set(qn("w:ascii"), family)
    rfonts.set(qn("w:hAnsi"), family)
    rfonts.set(qn("w:cs"), family)

    lig = OxmlElement("w14:ligatures")
    lig.set(qn("w14:val"), "standardContextual")
    rpr.append(lig)

    if lining_nums or oldstyle_nums:
        num_form = OxmlElement("w14:numForm")
        num_form.set(qn("w14:val"), "lining" if lining_nums else "oldstyle")
        rpr.append(num_form)
    if tabular_nums or proportional_nums:
        num_spacing = OxmlElement("w14:numSpacing")
        num_spacing.set(qn("w14:val"), "tabular" if tabular_nums else "proportional")
        rpr.append(num_spacing)


def set_paragraph_bottom_rule(paragraph, color_hex: str, padding_pt: float) -> None:
    ppr = paragraph._element.get_or_add_pPr()
    p_bdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "4")  # 0.5pt in eighths of a point
    bottom.set(qn("w:space"), str(max(1, int(round(padding_pt)))))
    bottom.set(qn("w:color"), color_hex.lstrip("#").upper())
    p_bdr.append(bottom)
    ppr.append(p_bdr)


def set_line_spacing(paragraph, line_pt: float) -> None:
    paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    paragraph.paragraph_format.line_spacing = Pt(line_pt)


def strip_table_borders(table) -> None:
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    if tbl_pr is None:
        tbl_pr = OxmlElement("w:tblPr")
        tbl.insert(0, tbl_pr)
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = OxmlElement(f"w:{edge}")
        element.set(qn("w:val"), "nil")
        borders.append(element)
    tbl_pr.append(borders)


def set_default_style(doc: Document, canon: dict) -> None:
    style = doc.styles["Normal"]
    style.font.name = canon["fonts"]["familyText"]
    style.font.size = Pt(canon["scale"]["s0Pt"])
    style.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    style.paragraph_format.line_spacing = Pt(canon["scale"]["uPt"])
    style.paragraph_format.space_before = Pt(0)
    style.paragraph_format.space_after = Pt(0)


def add_label(doc: Document, canon: dict, text: str, *, first: bool = False) -> None:
    s0 = canon["scale"]["s0Pt"]
    u = canon["scale"]["uPt"]
    tracking_pt = canon["tracking"]["labelEm"] * s0
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.keep_with_next = True
    paragraph.paragraph_format.space_before = Pt(0 if first else canon["rhythm"]["preSectionPt"] / 2)
    paragraph.paragraph_format.space_after = Pt(canon["rhythm"]["labelGapPt"])
    set_line_spacing(paragraph, u)
    # Cap-offset optical: pull label up like CSS negative margin.
    if canon["metrics"].get("capOffsetS0Pt"):
        paragraph.paragraph_format.space_before = Pt(
            max(0, (0 if first else canon["rhythm"]["preSectionPt"] / 2) - canon["metrics"]["capOffsetS0Pt"])
        )
    run = paragraph.add_run(text.upper())
    set_run_font(run, canon["fonts"]["familyText"], s0, spacing_pt=tracking_pt)
    set_paragraph_bottom_rule(paragraph, canon["optical"]["ruleColor"], canon["rhythm"]["preSectionPt"] / 2)


def add_name_block(doc: Document, canon: dict, profile: dict) -> None:
    name_pt = canon["scale"]["namePt"]
    name_leading = canon["scale"]["nameLeadingPt"]
    s0 = canon["scale"]["s0Pt"]
    u = canon["scale"]["uPt"]
    tracking_pt = canon["tracking"]["displayEm"] * name_pt

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_before = Pt(0)
    title.paragraph_format.space_after = Pt(0)
    set_line_spacing(title, name_leading)
    title_run = title.add_run(profile.get("name", ""))
    set_run_font(title_run, canon["fonts"]["familyTitle"], name_pt, spacing_pt=tracking_pt)

    contact = profile.get("contact") or {}
    bits = [contact.get("location"), contact.get("phone"), contact.get("email")]
    if contact.get("website", {}).get("href"):
        bits.append(contact["website"].get("label") or contact["website"]["href"])
    elif contact.get("url"):
        bits.append(str(contact["url"]).replace("https://", "").replace("http://", "").rstrip("/"))
    for prof in contact.get("profiles") or []:
        bits.append(prof.get("url") or prof.get("username") or prof.get("network") or "")
    bits = [bit for bit in bits if bit]
    if not bits:
        return

    contact_para = doc.add_paragraph()
    contact_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    contact_para.paragraph_format.space_before = Pt(canon["rhythm"]["nameContactGapPt"])
    contact_para.paragraph_format.space_after = Pt(canon["rhythm"].get("contactToBodyGapPt", 7))
    set_line_spacing(contact_para, u)
    contact_run = contact_para.add_run(NNBSP_PIPE.join(bits))
    set_run_font(
        contact_run,
        canon["fonts"]["familyText"],
        s0,
        lining_nums=True,
        tabular_nums=True,
    )


def add_body(doc: Document, canon: dict, text: str, *, space_after: float = 0) -> None:
    paragraph = doc.add_paragraph()
    set_line_spacing(paragraph, canon["scale"]["uPt"])
    paragraph.paragraph_format.space_after = Pt(space_after)
    run = paragraph.add_run(text)
    set_run_font(
        run,
        canon["fonts"]["familyText"],
        canon["scale"]["s0Pt"],
        oldstyle_nums=True,
        proportional_nums=True,
    )


def add_bullets(doc: Document, canon: dict, lines: list[str]) -> None:
    indent_pt = canon["bullet"]["indentPt"]
    hang_pt = canon["bullet"]["hangEm"] * canon["scale"]["s0Pt"]
    text_gap_pt = canon["bullet"]["textGapEm"] * canon["scale"]["s0Pt"]
    left_indent = max(0, indent_pt - hang_pt)
    for index, line in enumerate(lines):
        paragraph = doc.add_paragraph()
        paragraph.paragraph_format.left_indent = Pt(left_indent)
        paragraph.paragraph_format.first_line_indent = Pt(-hang_pt)
        paragraph.paragraph_format.space_before = Pt(0 if index == 0 else 0)
        paragraph.paragraph_format.space_after = Pt(canon["bullet"]["gapPt"])
        set_line_spacing(paragraph, canon["bullet"]["linePt"])
        marker = paragraph.add_run("\u2022")
        set_run_font(marker, canon["fonts"]["familyText"], canon["scale"]["s0Pt"])
        gap = paragraph.add_run("\u00a0" * max(1, int(round(text_gap_pt / (canon["scale"]["s0Pt"] * 0.25)))))
        # Prefer a fixed NBSP gap approximating textGapEm rather than raw space crush.
        if text_gap_pt > 0:
            gap.text = "\u00a0"
        body = paragraph.add_run(line)
        set_run_font(
            body,
            canon["fonts"]["familyText"],
            canon["scale"]["s0Pt"],
            oldstyle_nums=True,
            proportional_nums=True,
        )


def experience_parts(entry: dict) -> tuple[str, str]:
    order = entry.get("headerOrder") or "company-first"
    if order == "role-first":
        return entry.get("role") or "", entry.get("companyLine") or entry.get("company") or ""
    return entry.get("company") or "", entry.get("roleLine") or entry.get("role") or ""


def add_entry_header(
    doc: Document,
    canon: dict,
    *,
    primary: str,
    secondary: str = "",
    date: str = "",
    stacked: bool = False,
    date_italic: bool = False,
) -> None:
    s0 = canon["scale"]["s0Pt"]
    u = canon["scale"]["uPt"]
    date_width = canon["metrics"]["dateReservedPt"]
    text_width = mm_to_pt(canon["margins"]["textWidthMm"])
    lead_width = max(40, text_width - date_width - canon["scale"]["colGutterPt"])
    gap = canon["rhythm"].get("entryTitleGapPt", u / 2)

    rows = 2 if stacked and secondary else 1
    table = doc.add_table(rows=rows, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    strip_table_borders(table)

    for row in table.rows:
        row.cells[0].width = Pt(lead_width)
        row.cells[1].width = Pt(date_width)

    if stacked and secondary:
        primary_para = table.rows[0].cells[0].paragraphs[0]
        set_line_spacing(primary_para, u)
        primary_run = primary_para.add_run(primary)
        set_run_font(primary_run, canon["fonts"]["familyText"], s0, bold=True)

        secondary_para = table.rows[1].cells[0].paragraphs[0]
        set_line_spacing(secondary_para, u)
        secondary_run = secondary_para.add_run(secondary)
        set_run_font(secondary_run, canon["fonts"]["familyText"], s0, italic=True)

        if date:
            date_para = table.rows[0].cells[1].paragraphs[0]
            date_para.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            set_line_spacing(date_para, u)
            date_run = date_para.add_run(date)
            set_run_font(
                date_run,
                canon["fonts"]["familyText"],
                s0,
                italic=date_italic,
                lining_nums=True,
                tabular_nums=True,
            )
    else:
        lead = table.rows[0].cells[0].paragraphs[0]
        set_line_spacing(lead, u)
        primary_run = lead.add_run(primary)
        set_run_font(primary_run, canon["fonts"]["familyText"], s0, bold=True)
        if secondary:
            # Entry-title gap via NBSP cluster approximating column-gap.
            spacer = lead.add_run("\u00a0")
            set_run_font(spacer, canon["fonts"]["familyText"], s0)
            if gap > s0 * 0.35:
                lead.add_run("\u00a0")
            secondary_run = lead.add_run(secondary)
            set_run_font(secondary_run, canon["fonts"]["familyText"], s0, italic=True)

        if date:
            date_para = table.rows[0].cells[1].paragraphs[0]
            date_para.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            set_line_spacing(date_para, u)
            date_run = date_para.add_run(date)
            set_run_font(
                date_run,
                canon["fonts"]["familyText"],
                s0,
                italic=date_italic,
                lining_nums=True,
                tabular_nums=True,
            )

    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_before = Pt(0)
    spacer.paragraph_format.space_after = Pt(canon["rhythm"]["leadInGapPt"])
    spacer.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    spacer.paragraph_format.line_spacing = Pt(0.1)


def add_entry_gap(doc: Document, canon: dict) -> None:
    gap = canon["rhythm"].get("entryGapPt", 0)
    if gap <= 0:
        return
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(gap)
    paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    paragraph.paragraph_format.line_spacing = Pt(0.1)


def embed_fonts(docx_path: Path, canon: dict) -> None:
    faces = canon["fonts"]["docxFaces"]
    temp_path = docx_path.with_suffix(".tmp.docx")
    shutil.copyfile(docx_path, temp_path)

    with zipfile.ZipFile(temp_path, "r") as zin:
        archive = {name: zin.read(name) for name in zin.namelist()}

    rels_path = "word/_rels/fontTable.xml.rels"
    font_table_path = "word/fontTable.xml"
    content_types_path = "[Content_Types].xml"
    document_rels_path = "word/_rels/document.xml.rels"

    if rels_path not in archive:
        archive[rels_path] = (
            b'<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
            b'<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"></Relationships>'
        )
    if font_table_path not in archive:
        archive[font_table_path] = (
            b'<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
            b'<w:fonts xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"></w:fonts>'
        )

    rels_root = ET.fromstring(archive[rels_path])
    fonts_root = ET.fromstring(archive[font_table_path])
    content_root = ET.fromstring(archive[content_types_path])
    existing_rels = len(rels_root.findall(f"{{{REL_NS}}}Relationship"))

    for face in faces:
        family = face["family"]
        font_path = Path(face["path"])
        if not font_path.exists():
            raise FileNotFoundError(font_path)
        rel_id = f"rId{existing_rels + 1}"
        existing_rels += 1
        font_key = uuid.uuid4()
        key = font_key.bytes[::-1]
        payload = bytearray(font_path.read_bytes())
        for index in range(32):
            payload[index] ^= key[index % 16]
        name = font_path.stem + ".odttf"
        archive_name = f"word/fonts/{name}"
        archive[archive_name] = bytes(payload)
        override = ET.SubElement(content_root, f"{{{CT_NS}}}Override")
        override.set("PartName", f"/{archive_name}")
        override.set("ContentType", "application/vnd.openxmlformats-officedocument.obfuscatedFont")
        rel = ET.SubElement(rels_root, f"{{{REL_NS}}}Relationship")
        rel.set("Id", rel_id)
        rel.set("Type", "http://schemas.openxmlformats.org/officeDocument/2006/relationships/font")
        rel.set("Target", f"fonts/{name}")
        font_el = next((font for font in fonts_root if font.get(f"{{{W_NS}}}name") == family), None)
        if font_el is None:
            font_el = ET.SubElement(fonts_root, f"{{{W_NS}}}font")
            font_el.set(f"{{{W_NS}}}name", family)
        embed = ET.SubElement(font_el, f"{{{W_NS}}}embed{face['style']}")
        embed.set(f"{{{R_NS}}}id", rel_id)
        embed.set(f"{{{W_NS}}}fontKey", "{" + str(font_key).upper() + "}")

    archive[rels_path] = ET.tostring(rels_root, encoding="utf-8", xml_declaration=True)
    archive[font_table_path] = ET.tostring(fonts_root, encoding="utf-8", xml_declaration=True)
    archive[content_types_path] = ET.tostring(content_root, encoding="utf-8", xml_declaration=True)

    if document_rels_path in archive:
        doc_rels = ET.fromstring(archive[document_rels_path])
        has_font_table = any(
            rel.get("Type") == "http://schemas.openxmlformats.org/officeDocument/2006/relationships/fontTable"
            for rel in doc_rels.findall(f"{{{REL_NS}}}Relationship")
        )
        if not has_font_table:
            rel_ids = [rel.get("Id", "") for rel in doc_rels.findall(f"{{{REL_NS}}}Relationship")]
            next_id = max((int(rid[3:]) for rid in rel_ids if rid.startswith("rId")), default=0) + 1
            rel = ET.Element(f"{{{REL_NS}}}Relationship")
            rel.set("Id", f"rId{next_id}")
            rel.set("Type", "http://schemas.openxmlformats.org/officeDocument/2006/relationships/fontTable")
            rel.set("Target", "fontTable.xml")
            doc_rels.append(rel)
            archive[document_rels_path] = ET.tostring(doc_rels, encoding="utf-8", xml_declaration=True)
            override = ET.Element(f"{{{CT_NS}}}Override")
            override.set("PartName", "/word/fontTable.xml")
            override.set(
                "ContentType",
                "application/vnd.openxmlformats-officedocument.wordprocessingml.fontTable+xml",
            )
            content_root = ET.fromstring(archive[content_types_path])
            content_root.append(override)
            archive[content_types_path] = ET.tostring(content_root, encoding="utf-8", xml_declaration=True)

    # Ensure w14 namespace on document for OT figure features.
    document_xml = archive.get("word/document.xml")
    if document_xml:
        text = document_xml.decode("utf-8")
        if "xmlns:w14=" not in text and "w14:" in text:
            text = text.replace(
                "xmlns:w=",
                'xmlns:w14="http://schemas.microsoft.com/office/word/2010/wordml" xmlns:w=',
                1,
            )
            archive["word/document.xml"] = text.encode("utf-8")

    with zipfile.ZipFile(docx_path, "w", compression=zipfile.ZIP_DEFLATED) as zout:
        for name, data in archive.items():
            zout.writestr(name, data)
    temp_path.unlink(missing_ok=True)


def build_document(profile: dict, canon: dict) -> Document:
    doc = Document()
    section = doc.sections[0]
    section.page_height = Mm(canon["page"]["heightMm"])
    section.page_width = Mm(canon["page"]["widthMm"])
    section.top_margin = Mm(canon["margins"]["marginTopMm"])
    section.bottom_margin = Mm(canon["margins"]["marginBottomMm"])
    section.left_margin = Mm(canon["margins"]["marginXMm"])
    section.right_margin = Mm(canon["margins"]["marginXMm"])

    settings = doc.settings.element
    for flag in ("embedTrueTypeFonts", "embedSystemFonts"):
        settings.append(OxmlElement("w:" + flag))
    set_default_style(doc, canon)
    add_name_block(doc, canon, profile)

    if profile.get("summary"):
        add_label(doc, canon, profile.get("summaryLabel") or "Professional Summary", first=True)
        add_body(doc, canon, profile["summary"], space_after=canon["rhythm"].get("summaryTailGapPt", 3.5))

    if profile.get("experience"):
        add_label(doc, canon, "Professional Experience" if profile.get("isMasterCV") else "Relevant Experience")
        for index, entry in enumerate(profile["experience"]):
            primary, secondary = experience_parts(entry)
            add_entry_header(
                doc,
                canon,
                primary=primary,
                secondary=secondary,
                date=entry.get("date", ""),
                stacked=bool(entry.get("stacked")),
            )
            add_bullets(doc, canon, entry.get("highlights") or [])
            if index < len(profile["experience"]) - 1:
                add_entry_gap(doc, canon)

    if profile.get("projects"):
        add_label(doc, canon, "Projects" if profile.get("isMasterCV") else "Related Projects")
        for index, entry in enumerate(profile["projects"]):
            add_entry_header(
                doc,
                canon,
                primary=entry.get("name", ""),
                secondary=entry.get("engagementLabel") or "",
                date=entry.get("date") or "",
                stacked=bool(entry.get("stacked")),
            )
            if entry.get("description"):
                add_body(doc, canon, entry["description"])
            add_bullets(doc, canon, entry.get("highlights") or [])
            if index < len(profile["projects"]) - 1:
                add_entry_gap(doc, canon)

    if profile.get("education"):
        add_label(doc, canon, "Education")
        s0 = canon["scale"]["s0Pt"]
        after_colon_pt = canon["education"]["afterColonEm"] * s0
        colon_hang_pt = canon["education"]["colonHangEm"] * s0
        for index, entry in enumerate(profile["education"]):
            institution = entry.get("institution", "")
            location = entry.get("location", "")
            primary = institution if not location else f"{institution}{NNBSP_PIPE}{location}"
            add_entry_header(
                doc,
                canon,
                primary=primary,
                date=entry.get("date", ""),
                date_italic=True,
            )
            degree_bits = [bit for bit in [entry.get("school"), entry.get("studyType")] if bit]
            if degree_bits:
                add_body(doc, canon, NNBSP_PIPE.join(degree_bits))
            for label, key in [
                ("Double Major:", "area"),
                ("GPA:", "score"),
                ("Focus:", "summary"),
                ("Honors:", "honors"),
                ("Relevant Coursework:", "courses"),
            ]:
                value = entry.get(key)
                if key in ("honors", "courses") and isinstance(value, list):
                    value = "; ".join(value)
                if not value:
                    continue
                paragraph = doc.add_paragraph()
                paragraph.paragraph_format.left_indent = Pt(canon["education"]["detailInsetPt"])
                if colon_hang_pt:
                    paragraph.paragraph_format.first_line_indent = Pt(-colon_hang_pt)
                set_line_spacing(paragraph, canon["scale"]["uPt"])
                label_run = paragraph.add_run(label)
                set_run_font(
                    label_run,
                    canon["fonts"]["familyText"],
                    s0,
                    italic=True,
                    bold=(key == "score"),
                )
                # Approximate --edu-after-colon with NBSP padding.
                gap_run = paragraph.add_run("\u00a0" if after_colon_pt > 0 else " ")
                set_run_font(gap_run, canon["fonts"]["familyText"], s0)
                value_run = paragraph.add_run(str(value))
                set_run_font(
                    value_run,
                    canon["fonts"]["familyText"],
                    s0,
                    bold=(key == "score"),
                    oldstyle_nums=True,
                    proportional_nums=True,
                )
            if index < len(profile["education"]) - 1:
                add_entry_gap(doc, canon)

    additional = profile.get("additional") or {}
    if additional.get("visible") is False:
        return doc
    extra = []
    if profile.get("isMasterCV") and additional.get("skillMap"):
        for group in additional["skillMap"]:
            label = group.get("label") or group.get("name") or "Skills"
            extra.append((label, ", ".join(group.get("keywords") or [])))
    elif additional.get("skills"):
        extra.append(("Skills", ", ".join(additional["skills"])))
    if additional.get("certifications"):
        extra.append(("Certifications", ", ".join(additional["certifications"])))
    if additional.get("languages"):
        extra.append(("Languages", ", ".join(additional["languages"])))
    if additional.get("workAuthorization"):
        extra.append(("Work authorization", additional["workAuthorization"]))
    if additional.get("leadership"):
        extra.append(("Leadership", NNBSP_PIPE.join(additional["leadership"])))
    if extra:
        add_label(doc, canon, "Additional Information")
        for label, value in extra:
            paragraph = doc.add_paragraph()
            set_line_spacing(paragraph, canon["scale"]["uPt"])
            label_run = paragraph.add_run(f"{label}:")
            set_run_font(
                label_run,
                canon["fonts"]["familyText"],
                canon["scale"]["s0Pt"],
                bold=True,
                italic=True,
            )
            gap = paragraph.add_run("\u00a0")
            set_run_font(gap, canon["fonts"]["familyText"], canon["scale"]["s0Pt"])
            value_run = paragraph.add_run(value)
            set_run_font(
                value_run,
                canon["fonts"]["familyText"],
                canon["scale"]["s0Pt"],
                oldstyle_nums=True,
                proportional_nums=True,
            )

    return doc


def main() -> int:
    if len(sys.argv) != 4:
        print("usage: emit-docx.py <profile.json> <canon.json> <output.docx>", file=sys.stderr)
        return 2

    profile = load_json(Path(sys.argv[1]).resolve())
    canon = load_json(Path(sys.argv[2]).resolve())
    docx_path = Path(sys.argv[3]).resolve()
    staging = docx_path.with_suffix(".staging.docx")
    docx_path.parent.mkdir(parents=True, exist_ok=True)

    doc = build_document(profile, canon)
    doc.save(str(staging))
    embed_fonts(staging, canon)
    if docx_path.exists():
        docx_path.unlink()
    staging.replace(docx_path)
    print(f"docx:{docx_path}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
